import streamlit as st
import google.generativeai as genai
from PIL import Image
import json
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ฝัง API Key เรียบร้อย ไม่ต้องกรอกหน้าเว็บ
API_KEY_FIXED = "AQ.Ab8RN6IMrLQCS0hCIyr4yKaAY1oTn9vzmERBv8CbWc_qQE7DeA"
genai.configure(api_key=API_KEY_FIXED)

st.set_page_config(page_title="แปลงใบขอเบิกเป็น Excel", page_icon="📄", layout="centered")

st.title("📄 ระบบแปลงใบขอเบิก (Reservation) เป็น Excel")
st.write("อัปโหลดรูปภาพใบขอเบิก ระบบจะอ่านข้อมูลและสร้างไฟล์ Excel ให้คุณดาวน์โหลดได้ทันที")

uploaded_file = st.file_uploader("เลือกรูปภาพใบขอเบิก (JPG, PNG)", type=["jpg", "jpeg", "png"])

if uploaded_file is not None:
    image = Image.open(uploaded_file)
    st.image(image, caption="รูปภาพที่อัปโหลด", use_container_width=True)
    
    if st.button("🚀 เริ่มแปลงข้อมูลเป็น Excel", type="primary"):
        with st.spinner("กำลังให้ AI อ่านข้อมูลจากรูปภาพ รอสักครู่นะครับ..."):
            try:
                model = genai.GenerativeModel("gemini-3.8-flash")
                
                prompt = """
                จงอ่านข้อมูลจากภาพเอกสารใบขอเบิกนี้อย่างละเอียด แล้วสกัดข้อมูลออกมาเป็นรูปแบบ JSON 
                โดยโครงสร้างต้องเป็นดังนี้เท่านั้น (ห้ามใส่ข้อความอื่นนอกเหนือจาก JSON):
                {
                  "header": {
                    "reservation_no": "เลขที่ใบขอเบิก เช่น 1000240078",
                    "print_date": "วันที่พิมพ์",
                    "req_date": "วันที่ต้องการ",
                    "movement_type": "รหัสเคลื่อนไหว",
                    "receiving_plant": "ไปสถานที่ทำงาน",
                    "receiving_storage": "ไปสถานที่จัดเก็บ"
                  },
                  "items": [
                    {
                      "item": 1,
                      "material": "รหัสวัสดุ",
                      "description": "รายละเอียดสินค้า",
                      "quantity": 500.00,
                      "unit": "หน่วย",
                      "lot_no": "ล็อตสินค้า",
                      "plant": 1010,
                      "sloc": "SLoc",
                      "remark": "หมายเหตุ"
                    }
                  ]
                }
                แปลงค่า Quantity ให้เป็นตัวเลขทศนิยม (float) เสมอ
                """
                
                response = model.generate_content([image, prompt])
                clean_json = response.text.replace("```json", "").replace("```", "").strip()
                data = json.loads(clean_json)
                
                header = data.get("header", {})
                items = data.get("items", [])
                
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Reservation"
                ws.views.sheetView[0].showGridLines = True
                
                font_title = Font(name='Angsana New', size=16, bold=True)
                font_header = Font(name='Angsana New', size=13, bold=True, color="FFFFFF")
                font_data = Font(name='Angsana New', size=13)
                fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                fill_zebra = PatternFill(start_color="F9F9F9", end_color="F9F9F9", fill_type="solid")
                thin_border = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'), top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
                
                ws['A1'] = "ใบขอเบิก/ใบขอโอน (Reservation)"
                ws['A1'].font = font_title
                
                header_labels = {
                    "เลขที่ใบขอเบิก/Reservation": header.get("reservation_no", ""),
                    "วันที่พิมพ์/Print Date": header.get("print_date", ""),
                    "วันที่ต้องการ/Requirement Date": header.get("req_date", ""),
                    "รหัสเคลื่อนไหว/Movement Type": header.get("movement_type", ""),
                    "ไปสถานที่ทำงาน/Receiving Plant": header.get("receiving_plant", ""),
                    "ไปสถานที่จัดเก็บ/Receiving Storage": header.get("receiving_storage", "")
                }
                
                row_idx = 3
                for key, val in header_labels.items():
                    ws.cell(row=row_idx, column=1, value=key).font = Font(name='Angsana New', size=12, bold=True)
                    ws.cell(row=row_idx, column=2, value=val).font = Font(name='Angsana New', size=12)
                    row_idx += 1
                    
                row_idx += 1
                table_headers = ["Item", "Material", "Description", "Quantity", "Unit", "Lot No.", "Plant", "SLoc", "Remark"]
                for col_num, title in enumerate(table_headers, 1):
                    cell = ws.cell(row=row_idx, column=col_num, value=title)
                    cell.font = font_header
                    cell.fill = fill_header
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    
                start_table_row = row_idx
                row_idx += 1
                
                for item in items:
                    ws.cell(row=row_idx, column=1, value=item.get("item")).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=2, value=str(item.get("material", ""))).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=3, value=item.get("description", "")).alignment = Alignment(horizontal="left")
                    
                    qty_cell = ws.cell(row=row_idx, column=4, value=float(item.get("quantity", 0)))
                    qty_cell.number_format = '#,##0.000'
                    qty_cell.alignment = Alignment(horizontal="right")
                    
                    ws.cell(row=row_idx, column=5, value=item.get("unit", "")).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=6, value=item.get("lot_no", "")).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=7, value=item.get("plant", "")).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=8, value=item.get("sloc", "")).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=9, value=item.get("remark", "")).alignment = Alignment(horizontal="center")
                    
                    for col_num in range(1, 10):
                        c = ws.cell(row=row_idx, column=col_num)
                        c.font = font_data
                        c.border = thin_border
                        if row_idx % 2 == 1:
                            c.fill = fill_zebra
                    row_idx += 1
                    
                total_row = row_idx
                ws.cell(row=total_row, column=3, value="รวมทั้งหมด (Total Quantity)").font = Font(name='Angsana New', size=13, bold=True)
                ws.cell(row=total_row, column=3).alignment = Alignment(horizontal="right")
                
                sum_cell = ws.cell(row=total_row, column=4, value=f"=SUM(D{start_table_row+1}:D{total_row-1})")
                sum_cell.font = Font(name='Angsana New', size=13, bold=True)
                sum_cell.number_format = '#,##0.000'
                sum_cell.alignment = Alignment(horizontal="right")
                
                for col_num in range(1, 10):
                    c = ws.cell(row=total_row, column=col_num)
                    c.border = Border(top=Side(style='thin', color='000000'), bottom=Side(style='double', color='000000'))
                    
                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
                    
                output = io.BytesIO()
                wb.save(output)
                output.seek(0)
                
                st.success(f"🎉 แปลงข้อมูลสำเร็จ! พบรายการสินค้าทั้งหมด {len(items)} รายการ")
                
                file_name = f"Reservation_{header.get('reservation_no', 'Output')}.xlsx"
                st.download_button(
                    label="📥 คลิกดาวน์โหลดไฟล์ Excel",
                    data=output,
                    file_name=file_name,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
                
            except Exception as e:
                st.error(f"เกิดข้อผิดพลาดในการประมวลผล: {e}")